"""شاشة المخزون: عرض/تعديل الأدوية، الكميات، الأسعار، الباركود، واستيراد/تصدير إكسل."""
import re
from datetime import date, datetime, timedelta
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem,
    QPushButton, QLabel, QDialog, QFormLayout, QLineEdit,
    QDateEdit, QMessageBox, QHeaderView, QFileDialog, QAbstractItemView,
    QScrollArea, QFrame, QCheckBox, QComboBox, QProgressDialog
)
from PySide6.QtCore import Qt, QDate, QTimer, QThread, Signal
from PySide6.QtGui import QColor
from sqlalchemy import or_, func, case, and_
from sqlalchemy.orm import joinedload

from app.db.database import get_session
from app.db.models import Product, ProductUnit, Batch, StockMovement, InvoiceItem, Invoice, Supplier, PurchaseOrder, PurchaseOrderItem
from app.db.approved_helper import approved_clause
from app.db.settings_helper import get_usd_rate, get_expiry_alert_days
from app.ui.widgets import NumberLineEdit, disable_scroll, enable_touch_scroll, create_usd_price_row
from app.ui.icons import icon


def _read_date(date_edit):
    """يقرأ قيمة QDateEdit بشكل موثوق. المشكلة: QDateEdit أحيانًا ما يحدّث قيمته
    الداخلية فورًا وأنت تكتب بالكيبورد - يحتاج تفسير النص (interpretText) اللي
    عادة يصير تلقائيًا بس لو ضغطت Enter أو طلعت من الحقل (focus out). هذا يجبره
    يفسّر أي نص مكتوب حاليًا بالحقل قبل ما نقرأ القيمة، حتى لو ما ضغطت Enter."""
    date_edit.interpretText()
    return date_edit.date().toPython()


def _strips_from_cartons(qty_cartons, strips_per_carton):
    """يحوّل كمية بالباكيت (تقدر تكون كسرية - 0.25، 0.5... لأدوية غالية
    تُشترى بجزء من الباكيت) للكمية الفعلية بالشريط (الوحدة الأساس
    المخزّنة فعليًا بالمخزون - Batch.quantity_available عمود صحيح Integer،
    لأنه ماكو معنى فيزيائي لـ"نص شريط" بالمخزون).

    نقرّب لأقرب شريط كامل (round half up، مو int() المباشر اللي يقطع
    الكسر نزولًا للصفر دائمًا - كان هذا يسبب مشكلة حقيقية: باكيت فيه 3
    أشرطة وشريت ربعه (0.25) كان يطلع 0.75 شريط، و int() يقطعها بيصير 0
    شريط بالكامل، يعني تضيع القيمة كلها من الفاتورة بصمت بدون أي تنبيه!
    التقريب لأقرب شريط كامل أدق وأسلم من التصفير الصامت)."""
    return int(qty_cartons * strips_per_carton + 0.5)


def _fmt_qty(n):
    """يعرض الكمية بالباكيت كرقم صحيح لو مضبوطة بالضبط، وإلا برقم عشري واحد
    (مثلاً 12 باكيت كاملة تُعرض 12، بينما 12 باكيت وثلث شريط تُعرض 12.3)."""
    return f"{n:,.0f}" if abs(n - round(n)) < 0.01 else f"{n:,.1f}"


class AddProductDialog(QDialog):
    """
    إضافة دواء جديد.
    الشراء دائمًا بالباكيت/الكرتون، والبيع يصير بالشريط أو بالباكيت - فالحقول
    هنا مبنية على هذا الأساس: تدخل سعر شراء الباكيت وعدد الأشرطة بداخله،
    وسعر شراء الشريط ينحسب تلقائيًا (تقسيم سعر الباكيت على عدد الأشرطة)،
    وتدخل سعر بيع الشريط وسعر بيع الباكيت كل وحدة لحالها (بدون أي احتساب تلقائي
    خاطئ يفترض إن الباكيت = 10 أضعاف الشريط).
    """
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("إضافة دواء جديد")
        self.setWindowModality(Qt.ApplicationModal)
        self.setMinimumWidth(400)
        self.setMaximumHeight(600)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)

        scroll = QScrollArea()
        enable_touch_scroll(scroll)
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        outer.addWidget(scroll)

        form_widget = QWidget()
        layout = QFormLayout(form_widget)
        layout.setContentsMargins(16, 16, 16, 16)
        scroll.setWidget(form_widget)

        self.name_input = QLineEdit()
        self.generic_input = QLineEdit()
        self.manufacturer_input = QLineEdit()
        self.category_input = QLineEdit()
        self.barcode_input = QLineEdit()

        self.strips_per_carton_input = NumberLineEdit(placeholder="3")
        self.carton_cost_input = NumberLineEdit()
        self.strip_cost_preview = QLabel("سعر شراء الشريط: —")
        self.strip_cost_preview.setStyleSheet("color:#6B7280;font-size:12px;")

        self.sale_price_input = NumberLineEdit()
        self.carton_price_input = NumberLineEdit()

        self.qty_cartons_input = NumberLineEdit(decimals=2)
        self.qty_strips_preview = QLabel("الكمية بالشريط: —")
        self.qty_strips_preview.setStyleSheet("color:#6B7280;font-size:12px;")

        self.min_threshold_input = NumberLineEdit(placeholder="10")
        self.expiry_input = QDateEdit(QDate.currentDate().addYears(1))
        self.expiry_input.setCalendarPopup(True)
        disable_scroll(self.expiry_input)

        layout.addRow("الاسم التجاري:", self.name_input)
        layout.addRow("الاسم العلمي:", self.generic_input)
        layout.addRow("الشركة المصنعة:", self.manufacturer_input)
        layout.addRow("التصنيف:", self.category_input)
        layout.addRow("الباركود:", self.barcode_input)

        layout.addRow(QLabel("— الشراء (دائمًا بالباكيت) —"))
        layout.addRow("عدد الأشرطة بالباكيت:", self.strips_per_carton_input)
        layout.addRow("سعر شراء الباكيت:", self.carton_cost_input)
        usd_rate = get_usd_rate(get_session())
        usd_row, _usd_input = create_usd_price_row(self.carton_cost_input, usd_rate)
        layout.addRow("أو سعر شراء الباكيت بالدولار:", usd_row)
        layout.addRow(self.strip_cost_preview)

        layout.addRow(QLabel("— أسعار البيع —"))
        layout.addRow("سعر بيع الشريط:", self.sale_price_input)
        layout.addRow("سعر بيع الباكيت/العلبة:", self.carton_price_input)

        layout.addRow(QLabel("— الكمية المستلمة الآن —"))
        layout.addRow("الكمية (بالباكيت):", self.qty_cartons_input)
        layout.addRow(self.qty_strips_preview)

        layout.addRow("الحد الأدنى للتنبيه (بالشريط):", self.min_threshold_input)
        layout.addRow("تاريخ الصلاحية:", self.expiry_input)

        self.carton_cost_input.textChanged.connect(self._update_previews)
        self.strips_per_carton_input.textChanged.connect(self._update_previews)
        self.qty_cartons_input.textChanged.connect(self._update_previews)
        self.sale_price_input.textChanged.connect(self._update_previews)
        self.carton_price_input.textChanged.connect(self._update_previews)
        self._update_previews()

        # زر الحفظ ثابت خارج منطقة السكرول - دائمًا مرئي وتقدر تضغطه بأي وقت
        save_btn = QPushButton("حفظ")
        save_btn.setStyleSheet("background:qlineargradient(x1:0,y1:0,x2:0,y2:1,stop:0 #16A34A,stop:1 #15803D);color:white;border-radius:8px;padding:10px;font-weight:bold;")
        save_btn.clicked.connect(self._validate_and_accept)
        outer.addWidget(save_btn)

    def _strips_per_carton(self):
        # لو ما أدخلت سعر بيع الشريط، نعتبر المنتج "علبة وحدة بس" (زي شراب أو مرهم)
        # وعدد الأشرطة بالعلبة = 1 تلقائيًا - بدون ما يعتمد على الرقم اللي بالخانة
        if self.sale_price_input.value() == 0:
            return 1
        return int(self.strips_per_carton_input.value() or 3)

    def _update_previews(self):
        spc = self._strips_per_carton()
        carton_cost = self.carton_cost_input.value()
        strip_cost = (carton_cost / spc) if spc else 0
        if spc == 1:
            self.strip_cost_preview.setText("✓ منتج \"علبة وحدة بس\" (بدون أشرطة) - سعر الشريط والعلبة نفس الشي تلقائيًا")
        else:
            self.strip_cost_preview.setText(f"سعر شراء الشريط (تلقائي): {strip_cost:,.0f} د.ع")
        qty_cartons = self.qty_cartons_input.value()
        self.qty_strips_preview.setText(f"الكمية بالشريط (تلقائي): {_strips_from_cartons(qty_cartons, spc)}")

    def _validate_and_accept(self):
        carton_cost = self.carton_cost_input.value()
        carton_price = self.carton_price_input.value()
        sale_price = self.sale_price_input.value()
        if not self.name_input.text().strip():
            QMessageBox.warning(self, "تنبيه", "لازم تكتب الاسم التجاري للدواء.")
            return
        if carton_cost == 0 and carton_price == 0 and sale_price == 0:
            QMessageBox.warning(
                self, "تنبيه - ماكو أي سعر",
                "ما أدخلت أي سعر (لا شراء ولا بيع). أدخل سعر واحد على الأقل قبل الحفظ."
            )
            return
        self.accept()

    def get_data(self):
        spc = self._strips_per_carton()
        carton_cost = self.carton_cost_input.value()
        strip_cost = (carton_cost / spc) if spc else 0
        qty_cartons = self.qty_cartons_input.value()
        sale_price = self.sale_price_input.value()
        carton_price = self.carton_price_input.value()
        if spc == 1:
            # علبة وحدة بس: نوحّد السعرين (لو كتب وحد بس ينعبّي الثاني بنفس القيمة)
            sale_price = sale_price or carton_price
            carton_price = carton_price or sale_price
        return {
            "name": self.name_input.text().strip(),
            "generic_name": self.generic_input.text().strip(),
            "manufacturer": self.manufacturer_input.text().strip(),
            "category": self.category_input.text().strip(),
            "barcode": self.barcode_input.text().strip(),
            "strips_per_carton": spc,
            "carton_cost": carton_cost,
            "cost": strip_cost,  # تكلفة الشريط - محسوبة تلقائيًا
            "sale_price": sale_price,
            "wholesale_price": 0,  # ملغى من الواجهة - الاعتماد صار على سعر شراء الباكيت فقط
            "carton_price": carton_price,
            "qty_cartons": qty_cartons,
            "qty": _strips_from_cartons(qty_cartons, spc),  # الكمية الفعلية المخزّنة (بالشريط)
            "min_threshold": int(self.min_threshold_input.value() or 10),
            "expiry": _read_date(self.expiry_input),
        }


class EditProductDialog(QDialog):
    """تعديل بيانات دواء موجود - يشمل إمكانية إضافة كمية جديدة (دفعة) مباشرة."""
    def __init__(self, product, parent=None):
        super().__init__(parent)
        self.product = product
        self.setWindowTitle(f"تعديل دواء: {product.name}")
        self.setWindowModality(Qt.ApplicationModal)
        self.setMinimumWidth(400)
        self.setMaximumHeight(600)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)

        scroll = QScrollArea()
        enable_touch_scroll(scroll)
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        outer.addWidget(scroll)

        form_widget = QWidget()
        layout = QFormLayout(form_widget)
        layout.setContentsMargins(16, 16, 16, 16)
        scroll.setWidget(form_widget)

        self.name_input = QLineEdit(product.name)
        self.generic_input = QLineEdit(product.generic_name or "")
        self.manufacturer_input = QLineEdit(product.manufacturer or "")
        self.category_input = QLineEdit(product.category or "")
        self.barcode_input = QLineEdit(product.barcode or "")
        self.sale_price_input = NumberLineEdit()
        self.sale_price_input.set_value(product.sale_price)
        self.carton_price_input = NumberLineEdit()
        self.carton_price_input.set_value(product.carton_price)
        self.strips_per_carton_input = NumberLineEdit(placeholder="3")
        self.strips_per_carton_input.set_value(product.strips_per_carton or 3)
        self.carton_purchase_price_input = NumberLineEdit()
        self.carton_purchase_price_input.set_value(product.carton_purchase_price)
        self.main_strip_cost_preview = QLabel("سعر شراء الشريط: —")
        self.main_strip_cost_preview.setStyleSheet("color:#6B7280;font-size:12px;")
        self.min_threshold_input = NumberLineEdit(placeholder="10")
        self.min_threshold_input.set_value(product.min_stock_threshold)

        current_stock = sum(b.quantity_available for b in product.batches)
        self.add_qty_cartons_input = NumberLineEdit(decimals=2, placeholder="0")
        self.add_carton_cost_input = NumberLineEdit(placeholder="0")
        self.add_strip_cost_preview = QLabel("سعر شراء الشريط: —")
        self.add_strip_cost_preview.setStyleSheet("color:#6B7280;font-size:12px;")
        self.add_qty_strips_preview = QLabel("الكمية بالشريط: —")
        self.add_qty_strips_preview.setStyleSheet("color:#6B7280;font-size:12px;")
        self.add_expiry_input = QDateEdit(QDate.currentDate().addYears(1))
        self.add_expiry_input.setCalendarPopup(True)
        disable_scroll(self.add_expiry_input)

        layout.addRow("الاسم التجاري:", self.name_input)
        layout.addRow("الاسم العلمي:", self.generic_input)
        layout.addRow("الشركة المصنعة:", self.manufacturer_input)
        layout.addRow("التصنيف:", self.category_input)
        layout.addRow("الباركود:", self.barcode_input)
        layout.addRow("سعر بيع الشريط:", self.sale_price_input)
        layout.addRow("سعر بيع الباكيت/العلبة:", self.carton_price_input)
        layout.addRow("عدد الأشرطة بالباكيت:", self.strips_per_carton_input)
        layout.addRow("سعر شراء الباكيت (الحالي):", self.carton_purchase_price_input)
        usd_rate = get_usd_rate(get_session())
        usd_row, _usd_input = create_usd_price_row(self.carton_purchase_price_input, usd_rate)
        layout.addRow("أو سعر شراء الباكيت بالدولار:", usd_row)
        self.update_existing_batches_checkbox = QCheckBox(
            "حدّث سعر شراء كل الدفعات الحالية بالمخزون بهذا السعر الجديد أيضًا"
        )
        self.update_existing_batches_checkbox.setStyleSheet("color:#6B7280;font-size:12px;")
        layout.addRow(self.update_existing_batches_checkbox)
        layout.addRow(self.main_strip_cost_preview)
        layout.addRow("الحد الأدنى للتنبيه:", self.min_threshold_input)

        self.carton_purchase_price_input.textChanged.connect(self._update_main_preview)
        self.strips_per_carton_input.textChanged.connect(self._update_main_preview)
        self.sale_price_input.textChanged.connect(self._update_main_preview)
        self.carton_price_input.textChanged.connect(self._update_main_preview)
        self._update_main_preview()

        # --- دفعات المخزون الحالية: تقدر تصحح تاريخ الصلاحية مباشرة بدون ما تحتاج تضيف كمية جديدة ---
        batches_sep = QLabel("— دفعات المخزون الحالية (عدّل الكمية أو تاريخ الصلاحية مباشرة لو غلط) —")
        batches_sep.setStyleSheet("color:#6B7280;font-size:12px;margin-top:8px;")
        layout.addRow(batches_sep)

        self._batch_inputs = []  # [(batch_id, QDateEdit, NumberLineEdit, spc_for_batch), ...]
        active_batches = [b for b in product.batches if b.quantity_available > 0]
        if not active_batches:
            layout.addRow(QLabel("ماكو دفعات فيها مخزون حاليًا."))
        else:
            for b in sorted(active_batches, key=lambda x: (x.expiry_date or QDate.currentDate().toPython())):
                row_widget = QWidget()
                row_h = QHBoxLayout(row_widget)
                row_h.setContentsMargins(0, 0, 0, 0)
                # عدد الأشرطة بالباكيت وقت شراء هذي الدفعة بالذات (مو عدد
                # الأشرطة الحالي للدواء، اللي ممكن يكون تغيّر بعدها) - لو دفعة
                # قديمة ماعندها هذا المسجّل (carton_strips_per_carton=0)،
                # نرجع لعدد الأشرطة الحالي كتقريب أفضل من لا شي.
                spc_batch = b.carton_strips_per_carton or (product.strips_per_carton or 3)
                # decimals=2 لأن الكمية المتبقية بالمخزون (بالشريط) مب لازم
                # تكون مضاعف مضبوط لعدد الأشرطة بالباكيت (مثلاً بعد بيع جزء من
                # الدفعة بالشريط المفرد) - لازم نقدر نعرض/ندخل كسور باكيت.
                qty_input = NumberLineEdit(decimals=2)
                qty_input.set_value(round(b.quantity_available / spc_batch, 2) if spc_batch else b.quantity_available)
                qty_input.setMaximumWidth(80)
                row_h.addWidget(qty_input)
                row_h.addWidget(QLabel("باكيت - صلاحية:"))
                date_edit = QDateEdit(b.expiry_date or QDate.currentDate().addYears(1).toPython())
                date_edit.setCalendarPopup(True)
                disable_scroll(date_edit)
                row_h.addWidget(date_edit)
                # سعر شراء الباكيت الفعلي المُدخل وقت شراء هذي الدفعة - لو
                # مسجّل نعرضه مباشرة بدون أي إعادة حساب. لو دفعة قديمة قبل
                # هذا التسجيل (carton_purchase_price=0) نرجع للتقريب القديم
                # (سعر الشريط × عدد الأشرطة الحالي) كبديل وحيد متوفر.
                if b.carton_purchase_price:
                    carton_price_display = b.carton_purchase_price
                else:
                    carton_price_display = (b.purchase_price or 0) * spc_batch
                price_lbl = QLabel(f"سعر شراء الباكيت: {carton_price_display:,.0f}")
                price_lbl.setStyleSheet("color:#6B7280;font-size:11px;margin-right:8px;")
                row_h.addWidget(price_lbl)
                # هوية الدفعة للعرض: اسم المورد + رقم الوصل بس (بدون كلمة
                # "دفعة" أو "الحالية") - مختصرة بصيغة "الكمية (مورد-رقم)".
                # لو الدفعة عليها هدية/بونص (bonus_quantity) نضيفها مفصولة
                # بعد اسم المورد/الوصل - "الكمية الأصلية + عدد الهدية" بالضبط
                # الطلب - حتى تنعرض بكل مكان يذكر مورد/وصل/كمية الدفعة مع بعض.
                supplier_name = b.supplier.name if b.supplier else "بدون مورد"
                batch_label = f"{supplier_name}-{b.receipt_number}" if b.receipt_number else supplier_name
                if b.bonus_quantity:
                    original_strips = b.quantity_received - b.bonus_quantity
                    bonus_cartons = round(b.bonus_quantity / spc_batch, 2) if spc_batch else b.bonus_quantity
                    original_cartons = round(original_strips / spc_batch, 2) if spc_batch else original_strips
                    batch_label += f" - {original_cartons:g} + هدية {bonus_cartons:g} باكيت"
                # الكلمة الأولى بالليبل: تاريخ استلام الدفعة لو نقدر نحدده
                # (بدل كلمة "الكمية" العامة اللي كانت تتكرر بكل دفعة بدون
                # فايدة تمييزية - خانة الرقم جنبها أصلًا واضح إنها كمية).
                # نجرب 3 مصادر بالترتيب - يشتغل تلقائيًا حتى على دفعات
                # قديمة سابقة لهذا التعديل، بدون أي حاجة لتحديث يدوي لبياناتها:
                #  1) رقم الوصل المسجّل على الدفعة نفسها مباشرة (b.receipt_number)
                #     - يُسجّل تلقائيًا وقت كل استلام فعلي من شاشة المشتريات،
                #     بغض النظر عن صيغة batch_number الداخلية (الأدق - يغطي
                #     أي دفعة عندها رقم وصل حتى لو batch_number قديم/مختلف).
                #     نبحث عن فاتورة شراء بنفس رقم الوصل (ونفس المورد لو
                #     محدد، لتفادي تطابق عرضي لو مورد ثاني استخدم نفس الرقم).
                #  2) فاتورة الشراء المرتبطة عبر batch_number بصيغة
                #     "PO-{order_id}-{product_id}" - احتياط لو رقم الوصل
                #     فاضي لأي سبب.
                #  3) لو ماكو ولا وحدة من فوق، نرجع لتاريخ حركة "شراء"
                #     المسجّلة لهذي الدفعة بجدول حركات المخزون.
                # دفعات ثانية أصلًا ماكو فاتورة أو رقم وصل أو حركة شراء
                # مرتبطة بيها (استيراد أولي بالجملة، تعديل يدوي، إرجاع
                # مبيعات) - نرجع لكلمة "الكمية" كبديل واضح بدل تاريخ فاضي
                # أو مخترع.
                label_prefix = "الكمية"
                order_date_val = None
                matched_order = None

                if b.receipt_number:
                    po_query = get_session().query(PurchaseOrder).filter(
                        PurchaseOrder.receipt_number == b.receipt_number
                    )
                    if b.supplier_id:
                        po_query = po_query.filter(PurchaseOrder.supplier_id == b.supplier_id)
                    candidate_order = po_query.order_by(PurchaseOrder.order_date.desc()).first()
                    if candidate_order and candidate_order.order_date:
                        order_date_val = candidate_order.order_date
                        matched_order = candidate_order

                if not order_date_val:
                    order_match = re.match(r"^PO-(\d+)-\d+$", b.batch_number or "")
                    if order_match:
                        candidate_order = get_session().query(PurchaseOrder).get(int(order_match.group(1)))
                        if candidate_order and candidate_order.order_date:
                            order_date_val = candidate_order.order_date
                            matched_order = candidate_order

                if not order_date_val:
                    purchase_movement = (
                        get_session().query(StockMovement)
                        .filter(StockMovement.batch_id == b.id, StockMovement.movement_type == "شراء")
                        .order_by(StockMovement.moved_at)
                        .first()
                    )
                    if purchase_movement and purchase_movement.moved_at:
                        order_date_val = purchase_movement.moved_at

                if order_date_val:
                    label_prefix = order_date_val.strftime("%Y-%m-%d")

                # ⚠️ لو لقينا فاتورة شراء فعلية مرتبطة (matched_order)، نخلي
                # التاريخ ورقم الوصل رابط قابل للضغط يفتح نافذة تفاصيل تلك
                # الفاتورة مباشرة - بدل ما تدورين عليها يدويًا بسجل الفواتير.
                batch_row_label = QLabel()
                if matched_order:
                    batch_row_label.setText(
                        f'<a href="#" style="color:#166534;">{label_prefix} ({batch_label})</a>:'
                    )
                    batch_row_label.setCursor(Qt.PointingHandCursor)
                    batch_row_label.linkActivated.connect(
                        lambda _checked=False, o=matched_order: self._open_batch_order_detail(o)
                    )
                else:
                    batch_row_label.setText(f"{label_prefix} ({batch_label}):")
                layout.addRow(batch_row_label, row_widget)
                self._batch_inputs.append((b.id, date_edit, qty_input, spc_batch))

        sep = QLabel(f"— أو أضف كمية جديدة الآن، بالباكيت (المخزون الحالي: {current_stock} شريط) —")
        sep.setStyleSheet("color:#6B7280;font-size:12px;margin-top:8px;")
        layout.addRow(sep)
        # المورد (اختياري) + رقم الوصل وتاريخه - نفس الحقول المتوفرة بشاشة
        # "فاتورة شراء جديدة" وبإضافة الموبايل، حتى إضافة كمية من هنا تسجّل
        # فاتورة شراء فعلية بحساب المورد (تظهر بكشف حسابه بالموردين)، مو
        # دفعة "يتيمة" بدون مصدر معروف زي ما كان يصير قبل هذا التعديل.
        self.add_supplier_combo = QComboBox()
        disable_scroll(self.add_supplier_combo)
        self.add_supplier_combo.addItem("— بدون مورد محدد —", None)
        session = get_session()
        for s in session.query(Supplier).order_by(Supplier.name).all():
            self.add_supplier_combo.addItem(s.name, s.id)
        layout.addRow("المورد (اختياري):", self.add_supplier_combo)
        self.add_receipt_number_input = QLineEdit()
        self.add_receipt_number_input.setPlaceholderText("رقم وصل الاستلام من المورد (اختياري)...")
        layout.addRow("رقم الوصل:", self.add_receipt_number_input)
        self.add_receipt_date_input = QDateEdit(QDate.currentDate())
        self.add_receipt_date_input.setCalendarPopup(True)
        disable_scroll(self.add_receipt_date_input)
        layout.addRow("تاريخ الوصل:", self.add_receipt_date_input)
        layout.addRow("كمية تضاف (بالباكيت):", self.add_qty_cartons_input)
        layout.addRow("سعر شراء هذي الدفعة (لو مختلف):", self.add_carton_cost_input)
        restock_usd_row, _restock_usd_input = create_usd_price_row(self.add_carton_cost_input, usd_rate)
        layout.addRow("أو سعرها بالدولار:", restock_usd_row)
        layout.addRow(self.add_strip_cost_preview)
        layout.addRow(self.add_qty_strips_preview)
        layout.addRow("تاريخ صلاحيتها:", self.add_expiry_input)

        self.add_carton_cost_input.textChanged.connect(self._update_add_previews)
        self.strips_per_carton_input.textChanged.connect(self._update_add_previews)
        self.add_qty_cartons_input.textChanged.connect(self._update_add_previews)
        self.sale_price_input.textChanged.connect(self._update_add_previews)
        self._update_add_previews()

        waste_sep = QLabel("— تسجيل هدر/تلف (يخصم من المخزون) —")
        waste_sep.setStyleSheet("color:#6B7280;font-size:12px;margin-top:8px;")
        layout.addRow(waste_sep)
        self.waste_qty_input = NumberLineEdit(placeholder="0")
        self.waste_reason_input = QLineEdit()
        self.waste_reason_input.setPlaceholderText("السبب (اختياري): انتهاء صلاحية، كسر...")
        layout.addRow("كمية الهدر (بالشريط):", self.waste_qty_input)
        layout.addRow("السبب:", self.waste_reason_input)

        # زر الحفظ ثابت خارج منطقة السكرول - دائمًا مرئي وتقدر تضغطه بأي وقت
        save_btn = QPushButton("حفظ التعديلات")
        save_btn.setStyleSheet("background:qlineargradient(x1:0,y1:0,x2:0,y2:1,stop:0 #16A34A,stop:1 #15803D);color:white;border-radius:8px;padding:10px;font-weight:bold;")
        save_btn.clicked.connect(self._validate_and_accept)
        outer.addWidget(save_btn)

    def _open_batch_order_detail(self, order):
        """يفتح نافذة تفاصيل فاتورة الشراء المرتبطة بدفعة معينة - راجع
        رابط التاريخ/رقم الوصل جنب كل دفعة أعلاه."""
        from app.ui.purchase_view import PurchaseOrderDetailDialog
        dialog = PurchaseOrderDetailDialog(order, get_session(), self)
        dialog.exec()

    def _validate_and_accept(self):
        if not self.name_input.text().strip():
            QMessageBox.warning(self, "تنبيه", "لازم تكتب الاسم التجاري للدواء.")
            return
        carton_cost = self.carton_purchase_price_input.value() or self.add_carton_cost_input.value()
        if (carton_cost == 0 and self.carton_price_input.value() == 0
                and self.sale_price_input.value() == 0):
            QMessageBox.warning(
                self, "تنبيه - ماكو أي سعر",
                "ما أدخلت أي سعر (لا شراء ولا بيع). أدخل سعر واحد على الأقل قبل الحفظ."
            )
            return
        self.accept()

    def _strips_per_carton(self):
        # نفس منطق نافذة الإضافة: تجاهل سعر بيع الشريط = منتج "علبة وحدة بس"
        if self.sale_price_input.value() == 0:
            return 1
        return int(self.strips_per_carton_input.value() or 3)

    def _update_main_preview(self):
        spc = self._strips_per_carton()
        carton_cost = self.carton_purchase_price_input.value()
        strip_cost = (carton_cost / spc) if spc else 0
        if spc == 1:
            self.main_strip_cost_preview.setText("✓ منتج \"علبة وحدة بس\" - سعر الشريط والعلبة نفس الشي تلقائيًا")
        else:
            self.main_strip_cost_preview.setText(f"سعر شراء الشريط (تلقائي): {strip_cost:,.0f} د.ع")

    def _update_add_previews(self):
        spc = self._strips_per_carton()
        carton_cost = self.add_carton_cost_input.value()
        strip_cost = (carton_cost / spc) if spc else 0
        self.add_strip_cost_preview.setText(f"سعر شراء الشريط (تلقائي): {strip_cost:,.0f} د.ع")
        qty_cartons = self.add_qty_cartons_input.value()
        self.add_qty_strips_preview.setText(f"الكمية بالشريط (تلقائي): {_strips_from_cartons(qty_cartons, spc)}")

    def get_data(self):
        spc = self._strips_per_carton()
        general_carton_cost = self.carton_purchase_price_input.value()
        qty_cartons = self.add_qty_cartons_input.value()
        # لو ما حدد سعر شراء خاص بهذي الدفعة، نستخدم "سعر شراء الباكيت (الحالي)" العام
        restock_carton_cost = self.add_carton_cost_input.value() or general_carton_cost
        strip_cost = (restock_carton_cost / spc) if spc else 0
        # الحقل العام يتحدث تلقائيًا لآخر سعر شراء فعلي لو صار فيه إضافة كمية جديدة
        new_reference_carton_cost = restock_carton_cost if qty_cartons > 0 else general_carton_cost
        sale_price = self.sale_price_input.value()
        carton_price = self.carton_price_input.value()
        if spc == 1:
            sale_price = sale_price or carton_price
            carton_price = carton_price or sale_price
        return {
            "name": self.name_input.text().strip(),
            "generic_name": self.generic_input.text().strip(),
            "manufacturer": self.manufacturer_input.text().strip(),
            "category": self.category_input.text().strip(),
            "barcode": self.barcode_input.text().strip(),
            "sale_price": sale_price,
            "carton_price": carton_price,
            "strips_per_carton": spc,
            "carton_purchase_price": new_reference_carton_cost,
            "update_existing_batches": self.update_existing_batches_checkbox.isChecked(),
            "min_threshold": int(self.min_threshold_input.value() or 10),
            "add_qty": _strips_from_cartons(qty_cartons, spc),  # الكمية الفعلية بالشريط
            "add_cost": strip_cost,             # تكلفة الشريط - محسوبة تلقائيًا من سعر الباكيت
            "add_carton_cost": restock_carton_cost,  # سعر شراء الباكيت الفعلي المُدخل لهذي الدفعة الجديدة
            "add_spc": spc,
            "add_expiry": _read_date(self.add_expiry_input),
            "add_supplier_id": self.add_supplier_combo.currentData(),
            "add_receipt_number": self.add_receipt_number_input.text().strip() or None,
            "add_receipt_date": _read_date(self.add_receipt_date_input),
            "waste_qty": int(self.waste_qty_input.value()),
            "waste_reason": self.waste_reason_input.text().strip(),
            "batch_updates": [
                # qty_input بهذا الحقل مُدخل بالباكيت (راجع بناء الصف بالأعلى) -
                # نحوّله لشريط هنا (الوحدة الفعلية المخزّنة) باستخدام عدد
                # الأشرطة الخاص بهذي الدفعة بالذات، مو عدد الأشرطة الحالي للدواء.
                (batch_id, _read_date(date_edit), round(qty_input.value() * batch_spc), batch_spc)
                for batch_id, date_edit, qty_input, batch_spc in self._batch_inputs
            ],
        }


class _ImportExcelWorker(QThread):
    """يشغّل استيراد ملف الإكسل كامل بخيط منفصل عن خيط الواجهة - قبل هذا
    الإصلاح، استيراد آلاف الأدوية كان يجمّد الواجهة تمامًا (يظهر "لا يستجيب"
    بويندوز) لأن كل استعلامات قاعدة البيانات كانت تشتغل على نفس الخيط اللي
    يرسم الشاشة، فويندوز يعتبرها متعلّقة حتى لو شغّالة فعليًا بالخلفية.
    يفتح جلسة (session) خاصة فيه بنفسه (get_session() تعتمد على
    scoped_session المرتبطة بالخيط الحالي - فتح الجلسة داخل run() هنا يضمن
    جلسة منفصلة آمنة عن جلسة الواجهة الرئيسية، بدل ما نشارك نفس الجلسة بين
    خيطين بنفس الوقت وهذا غير آمن إطلاقًا)."""

    progress = Signal(int, int)  # (الصف الحالي، إجمالي صفوف الملف)
    finished_ok = Signal(dict)   # ملخص النتيجة (انظر النتيجة بالأسفل)
    finished_error = Signal(str)

    def __init__(self, path):
        super().__init__()
        self.path = path
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def run(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.finished_error.emit("لازم تنصب openpyxl: pip install openpyxl")
            return

        session = get_session()
        try:
            wb = load_workbook(self.path, read_only=True, data_only=True)
            ws = wb.active
            total_rows = ws.max_row - 1 if ws.max_row and ws.max_row > 1 else 0

            count = 0
            inserted = 0
            updated = 0
            total_rows_seen = 0
            errors = []
            commit_failed = False
            cancelled = False

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if self._cancelled:
                    cancelled = True
                    break
                if row_num % 25 == 0 or row_num == total_rows + 1:
                    self.progress.emit(row_num - 1, total_rows)
                if not row or not row[0]:
                    continue
                total_rows_seen += 1
                try:
                    with session.begin_nested():
                        (name, generic, manufacturer, category, barcode, qty,
                         carton_purchase, strip_purchase, carton_sale, strip_sale, packing) = (
                            list(row) + [None] * 11
                        )[:11]
                        existing = session.query(Product).filter_by(name=name, is_active=True).first()
                        is_new = existing is None
                        if existing:
                            product = existing
                        else:
                            product = Product(name=name, is_custom=True, base_unit="شريط")
                            session.add(product)
                            session.flush()
                            for unit_name, factor in [("شريط", 1), ("علبة", product.strips_per_carton or 3)]:
                                session.add(ProductUnit(product_id=product.id, unit_name=unit_name, conversion_factor=factor))
                        product.generic_name = generic
                        product.manufacturer = manufacturer
                        product.category = category
                        product.barcode = str(barcode) if barcode else product.barcode

                        packing_val = self._parse_number(packing, as_int=True) or product.strips_per_carton or 1
                        product.strips_per_carton = packing_val
                        product.carton_purchase_price = self._parse_number(carton_purchase)
                        product.carton_price = self._parse_number(carton_sale)
                        product.sale_price = self._parse_number(strip_sale)

                        qty_val = self._parse_number(qty, as_int=True)
                        if qty_val:
                            spc = product.strips_per_carton or 1
                            strip_cost = (product.carton_purchase_price / spc) if spc else 0
                            session.add(Batch(
                                product_id=product.id, batch_number="IMPORT",
                                expiry_date=date.today() + timedelta(days=365),
                                purchase_price=strip_cost, quantity_received=qty_val, quantity_available=qty_val,
                                carton_purchase_price=product.carton_purchase_price, carton_strips_per_carton=spc,
                            ))
                    count += 1
                    if is_new:
                        inserted += 1
                    else:
                        updated += 1
                    if count % 500 == 0:
                        try:
                            session.commit()
                        except Exception as commit_err:
                            errors.append(f"فشل الحفظ عند الصف {row_num}: {commit_err}")
                            commit_failed = True
                            break
                except Exception as e:
                    errors.append(f"صف {row_num} ({row[0] if row and row[0] else '؟'}): {e}")

            if not commit_failed:
                try:
                    session.commit()
                except Exception as commit_err:
                    errors.append(f"فشل الحفظ النهائي: {commit_err}")
                    commit_failed = True

            self.finished_ok.emit({
                "count": count, "inserted": inserted, "updated": updated,
                "total_rows_seen": total_rows_seen, "errors": errors,
                "commit_failed": commit_failed, "cancelled": cancelled,
            })
        except Exception as e:
            self.finished_error.emit(str(e))
        finally:
            session.close()

    @staticmethod
    def _parse_number(value, as_int=False, default=0):
        if value is None:
            return default
        if isinstance(value, (int, float)):
            return int(value) if as_int else float(value)
        text = str(value).strip().replace(",", "")
        match = re.search(r"-?\d+(\.\d+)?", text)
        if not match:
            return default
        num = float(match.group())
        return int(num) if as_int else num


class InventoryView(QWidget):
    PAGE_SIZE = 100

    def showEvent(self, event):
        """يشتغل تلقائيًا كل مرة تصير صفحة المخزون مرئية - سواء أول فتح
        للبرنامج أو الرجوع لها من صفحة ثانية بشريط التنقل. نحط مؤشر الكتابة
        مباشرة بمربع البحث/الباركود بدون ما يحتاج المستخدم يضغط عليه يدويًا
        أول - عشان يقدر يمسح الباركود أو يكتب اسم الدواء على طول."""
        super().showEvent(event)
        self.barcode_scan_input.setFocus()

    def __init__(self, parent=None):
        super().__init__(parent)
        self.session = get_session()
        self._current_page = 1
        self._total_pages = 1
        layout = QVBoxLayout(self)

        header = QHBoxLayout()
        title = QLabel("المخزون")
        title.setStyleSheet("font-size:18px;font-weight:bold;")
        add_btn = QPushButton("+ إضافة دواء جديد")
        add_btn.setStyleSheet("background:qlineargradient(x1:0,y1:0,x2:0,y2:1,stop:0 #16A34A,stop:1 #15803D);color:white;border-radius:8px;padding:8px 14px;")
        add_btn.clicked.connect(self.add_product)
        # ملاحظة: زر "تصدير Excel" انشال من الواجهة بناءً على طلب صاحب
        # البرنامج (حماية بيانات المخزون من التصدير/التسريب) - الدالة
        # export_excel() نفسها خليناها بالكود (تحت) بس بلا أي زر يستدعيها،
        # حتى لو احتاج المطوّر يرجّعها بالمستقبل يعرف وين تكون.
        import_btn = QPushButton("استيراد من Excel")
        import_btn.setStyleSheet("background:#F1F5F9;border:1px solid #E5E7EB;border-radius:8px;padding:8px 14px;")
        import_btn.clicked.connect(self.import_excel)
        header.addWidget(title)
        header.addStretch()
        header.addWidget(import_btn)
        header.addWidget(add_btn)
        layout.addLayout(header)

        barcode_row = QHBoxLayout()
        self.barcode_scan_input = QLineEdit()
        self.barcode_scan_input.setPlaceholderText("ابحث بالاسم أو امسح الباركود هنا - يفتح تلقائيًا لو طابق باركود بالضبط")
        self.barcode_scan_input.returnPressed.connect(self.handle_barcode_scan)
        self.barcode_scan_input.textChanged.connect(self._on_barcode_typed)
        barcode_row.addWidget(self.barcode_scan_input)

        self.show_archived_check = QCheckBox("إظهار الأدوية المؤرشفة")
        self.show_archived_check.stateChanged.connect(lambda _: self.refresh(self.barcode_scan_input.text()))
        barcode_row.addWidget(self.show_archived_check)
        layout.addLayout(barcode_row)

        # فلتر "الأدوية المعتمدة فقط" - يخفي الأدوية اللي مجرد اسم وباركود
        # من الكتالوج الجاهز (starter_catalog.xlsx) ولم تُشترَ/تُستخدم أبدًا
        # بالصيدلية فعليًا (راجع app/db/approved_helper.py لتعريف "معتمد").
        # زر تفعيل/تطفيء بسيط تحت خانة "إظهار الأدوية المؤرشفة" مباشرة.
        approved_row = QHBoxLayout()
        self.approved_only_check = QCheckBox("عرض الأدوية المعتمدة فقط")
        self.approved_only_check.setToolTip(
            "يخفي الأدوية الموجودة بالكتالوج الجاهز كاسم وباركود بس ولم يتم شراؤها\n"
            "أو التعامل معها فعليًا بالصيدلية - يبقي فقط الأدوية اللي إلها أسعار\n"
            "وكميات حقيقية ومُعتمدة من الصيدلاني."
        )
        self.approved_only_check.stateChanged.connect(lambda _: self.refresh(self.barcode_scan_input.text()))
        approved_row.addWidget(self.approved_only_check)
        approved_row.addStretch()
        layout.addLayout(approved_row)

        barcode_row2 = QHBoxLayout()
        # فلترة حسب الحالة - تُستخدم يدويًا من هنا، أو تلقائيًا لما توصل من
        # زر "عرض الكل" بالرئيسية أو التقارير (عبر set_status_filter).
        barcode_row2.addWidget(QLabel("الحالة:"))
        self.status_filter_combo = QComboBox()
        self.status_filter_combo.addItems(["الكل", "نقص", "منخفض", "منتهي", "قريب الانتهاء", "راكد"])
        self.status_filter_combo.currentTextChanged.connect(lambda _: self.refresh(self.barcode_scan_input.text()))
        barcode_row2.addWidget(self.status_filter_combo)
        barcode_row2.addStretch()
        layout.addLayout(barcode_row2)

        self.table = QTableWidget()
        enable_touch_scroll(self.table)
        self.table.setAlternatingRowColors(True)
        self.table.setColumnCount(15)
        self.table.setHorizontalHeaderLabels([
            "الاسم التجاري", "الاسم العلمي", "الشركة", "التصنيف", "الباركود",
            "شراء\nالباكيت", "بيع\nالباكيت", "شراء\nالشريط", "بيع\nالشريط",
            "كمية\nالباكيت", "كمية\nالشريط",
            "أقرب\nصلاحية", "الحالة", "", ""
        ])
        # ملاحظة وضوح/حجم: الجدول كان يبين عريض والكلمات بالهيدر مو واضحة -
        # صغّرنا الخط وقلّلنا الحشوة الداخلية لكل خلية، وقصّرنا نصوص العناوين
        # (مثلاً "سعر شراء الباكيت" صارت "شراء الباكيت" بسطرين) مع تفعيل
        # التفاف النص بالهيدر عشان يبين كامل وواضح بدل ما ينقطع أو يتلخبط.
        self.table.setStyleSheet(
            "QTableWidget{font-size:12px;}"
            "QTableWidget::item{padding:3px 4px;}"
            "QHeaderView::section{font-size:11px;font-weight:700;padding:4px 3px;}"
        )
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setStretchLastSection(False)
        header.setSectionResizeMode(0, QHeaderView.Stretch)
        header.setSectionResizeMode(1, QHeaderView.Stretch)
        header.setMinimumSectionSize(50)
        try:
            header.setWordWrap(True)
        except AttributeError:
            pass  # نسخ Qt قديمة ما تدعم التفاف نص الهيدر - تبقى الأعمدة شغالة بدونه
        header.setMinimumHeight(40)
        fixed_widths = {
            2: 85, 3: 90, 4: 90, 5: 62, 6: 62, 7: 62, 8: 62,
            9: 58, 10: 58, 11: 72, 12: 62, 13: 52, 14: 52,
        }
        for col, w in fixed_widths.items():
            self.table.setColumnWidth(col, w)
        self.table.setEditTriggers(QAbstractItemView.DoubleClicked)
        self.table.itemChanged.connect(self._on_item_changed)
        layout.addWidget(self.table)

        # تقسيم لصفحات (Pagination): مع مخزون كبير (آلاف الأدوية)، كان الجدول
        # يبني صف Qt لكل دواء مطابق دفعة وحدة بكل تحديث/بحث - يخلي فتح
        # المخزون أو الفلترة بطيء بشكل ملحوظ. الحين نعرض 100 دواء بالصفحة
        # بس، ولو ماكو فلتر حالة نطبّق حتى تقسيم الصفحات على استعلام قاعدة
        # البيانات نفسه (LIMIT/OFFSET) بدل ما نجيب كل الصفوف ونقصّها ببايثون.
        pagination_row = QHBoxLayout()
        self.prev_page_btn = QPushButton("◀ السابق")
        self.prev_page_btn.setStyleSheet("background:#F1F5F9;border:1px solid #E5E7EB;border-radius:8px;padding:6px 14px;")
        self.prev_page_btn.clicked.connect(self._go_prev_page)
        self.page_info_label = QLabel("")
        self.page_info_label.setAlignment(Qt.AlignCenter)
        self.page_info_label.setStyleSheet("color:#374151;font-weight:600;")
        self.next_page_btn = QPushButton("التالي ▶")
        self.next_page_btn.setStyleSheet("background:#F1F5F9;border:1px solid #E5E7EB;border-radius:8px;padding:6px 14px;")
        self.next_page_btn.clicked.connect(self._go_next_page)
        pagination_row.addWidget(self.prev_page_btn)
        pagination_row.addWidget(self.page_info_label, stretch=1)
        pagination_row.addWidget(self.next_page_btn)
        layout.addLayout(pagination_row)

        self._loading = False
        self.refresh()

        # مؤقت تأخير بسيط: يفتح تلقائيًا بعد وقفة قصيرة بالكتابة (يشتغل مع الماسح الضوئي
        # حتى لو ما يرسل Enter بنهاية القراءة)
        self._barcode_timer = QTimer(self)
        self._barcode_timer.setSingleShot(True)
        self._barcode_timer.timeout.connect(self.handle_barcode_scan)

    def _on_barcode_typed(self, text):
        # بحث حي بالاسم/الباركود يفلتر الجدول فورًا وأنت تكتب
        self.refresh(search=text)
        # وبنفس الوقت، لو طابق النص باركود بالضبط (زي قارئ الباركود اللي يرسل النص كامل
        # دفعة وحدة) نفتح الدواء تلقائيًا بعد وقفة قصيرة بدون ما تحتاج تضغط Enter
        self._barcode_timer.stop()
        if len(text.strip()) >= 6:
            self._barcode_timer.start(350)

    def _go_prev_page(self):
        if self._current_page > 1:
            self._current_page -= 1
            self.refresh(self.barcode_scan_input.text(), reset_page=False)

    def _go_next_page(self):
        if self._current_page < self._total_pages:
            self._current_page += 1
            self.refresh(self.barcode_scan_input.text(), reset_page=False)

    def set_status_filter(self, status):
        """يُستدعى من main_window لما توصل من زر "عرض الكل" بالرئيسية أو
        التقارير - يختار نفس الحالة بقائمة الفلترة تلقائيًا ويحدّث الجدول،
        عشان المستخدم يوصل مباشرة لنفس المجموعة اللي كان يشوفها بدون ما
        يفلتر يدويًا من جديد."""
        if status and self.status_filter_combo.findText(status) >= 0:
            self.status_filter_combo.setCurrentText(status)
        else:
            self.refresh(self.barcode_scan_input.text())

    def _stagnant_product_ids(self):
        """أدوية إلها مخزون فعلي (غير منتهي) بس ما انباعت من 6 أشهر فما فوق -
        نفس تعريف "الراكدة" المستخدم بشاشة التقارير بالضبط، محسوب هنا بشكل
        منفصل ويُستدعى بس لما فلتر "راكد" مختار (تجنبًا لاستعلام إضافي غير
        ضروري بكل تحديث عادي للجدول).

        ⚠️ إصلاح أداء (راجع تقرير التدقيق، بند 14.5): كانت هذي الدالة تجيب
        *كل* منتج فعّال مع كل دفعاته كاملة (joinedload(Product.batches).all())
        بس عشان تحسب "المخزون القابل للاستخدام" (الدفعات غير المنتهية)
        بحلقة بايثون. الحين نفس التعريف بالضبط، بس محسوب بتجميع SQL
        (SUM مشروط بـ CASE WHEN مع GROUP BY) بدل تحميل كل كائنات
        المنتجات/الدفعات."""
        today = date.today()
        cutoff = datetime.now() - timedelta(days=180)
        last_sale_rows = (
            self.session.query(Batch.product_id, func.max(Invoice.invoice_date))
            .join(InvoiceItem, InvoiceItem.batch_id == Batch.id)
            .join(Invoice, InvoiceItem.invoice_id == Invoice.id)
            .filter(Invoice.payment_status != "مرتجع")
            .group_by(Batch.product_id)
            .all()
        )
        last_sale_by_product = {pid: last_date for pid, last_date in last_sale_rows}

        # "دفعة قابلة للاستخدام" = ماعندها تاريخ صلاحية أصلًا أو صلاحيتها
        # لسا ما انتهت (نفس شرط: not (b.expiry_date and b.expiry_date <=
        # today) بالضبط) - نجمع كمياتها بس (else 0) بدل تحميل كل الدفعات.
        usable_stock_rows = (
            self.session.query(
                Batch.product_id,
                func.coalesce(
                    func.sum(
                        case(
                            (or_(Batch.expiry_date.is_(None), Batch.expiry_date > today), Batch.quantity_available),
                            else_=0,
                        )
                    ),
                    0,
                ),
            )
            .join(Product, Product.id == Batch.product_id)
            .filter(Product.is_active == True)
            .group_by(Batch.product_id)
            .all()
        )

        ids = set()
        for pid, usable_stock in usable_stock_rows:
            if usable_stock <= 0:
                continue
            last_sale = last_sale_by_product.get(pid)
            if last_sale is None or last_sale < cutoff:
                ids.add(pid)
        return ids

    def handle_barcode_scan(self):
        code = self.barcode_scan_input.text().strip()
        if not code:
            return
        product = self.session.query(Product).filter_by(barcode=code).first()
        if product:
            self.barcode_scan_input.clear()
            self.edit_product(product)
        else:
            # ماكو تطابق باركود - نخليها بس فلترة بالاسم، ما نفتح شاشة "دواء جديد" تلقائيًا
            # (نفتحها فقط لو ضغط Enter بنفسه وهو متأكد إنه دواء جديد فعلاً)
            existing_by_name = self.session.query(Product).filter(Product.name.ilike(code)).first()
            if not existing_by_name:
                reply = QMessageBox.question(
                    self, "دواء جديد؟",
                    f"ماكو دواء بهذا الاسم/الباركود ({code}). تريد تضيفه كدواء جديد؟",
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.No
                )
                if reply == QMessageBox.Yes:
                    self.barcode_scan_input.clear()
                    dialog = AddProductDialog(self)
                    if len(code) >= 6 and code.isdigit():
                        dialog.barcode_input.setText(code)
                    else:
                        dialog.name_input.setText(code)
                    self._save_new_product(dialog)

    def refresh(self, search="", reset_page=True):
        if reset_page:
            self._current_page = 1
        self._loading = True
        # ⚠️ إصلاح خلل نادر (كميات تُعرض 0 رغم وجودها فعليًا بقاعدة البيانات):
        # self.session جلسة SQLAlchemy طويلة العمر (scoped_session بنفس خيط
        # الواجهة، تضل مفتوحة من فتح البرنامج لين إغلاقه). لو دواء معيّن
        # انحمّل مرة بذاكرة هذي الجلسة (identity map)، وبعدين انضاف له مخزون
        # من خيط ثاني تمامًا - خادم الموبايل (app/mobile_server.py، endpoints
        # مسوّاة sync def فتشتغل بخيط threadpool منفصل) أو استيراد الإكسل
        # (QThread منفصل بنفس هذا الملف) - جلسة هذي الشاشة ما "تعرف" بالتغيير
        # تلقائيًا، وتضل تعرض القيمة القديمة المحمّلة سابقًا لحد ما يصير أي
        # commit ثاني بنفس خيط الواجهة (يفرّغ الكاش تلقائيًا) أو يعاد فتح
        # البرنامج. expire_all() تفرّغ كاش هذي الجلسة بس (بدون ما تمسح أي
        # بيانات فعلية) فيرجع أي استعلام بعدها يجيب القيم الحقيقية الحالية
        # من القرص - نفس فكرة session.refresh(card.product) المستخدمة أصلًا
        # بـ pos_view.py لنفس فئة المشكلة بالضبط.
        self.session.expire_all()
        # يُقرأ مرة وحدة هنا بس (مو داخل _status_for لكل صف) حتى ما نفتح
        # استعلام إعدادات إضافي لكل دواء بالجدول - يُستخدم كإعداد "قرب
        # انتهاء الصلاحية" القابل للتعديل من الإعدادات بدل الرقم الثابت 90.
        self._expiry_alert_days = get_expiry_alert_days(self.session)
        # joinedload(Product.batches) يجيب الدفعات مع المنتجات باستعلام واحد
        # بدل ما تفتح استعلام Batch منفصل لكل صف بالجدول (N+1) - هذا كان يخلي
        # فتح شاشة المخزون أبطأ كل ما زاد عدد الأدوية.
        query = self.session.query(Product).options(joinedload(Product.batches))
        if not self.show_archived_check.isChecked():
            query = query.filter(Product.is_active == True)
        if self.approved_only_check.isChecked():
            query = query.filter(approved_clause(self.session))
        search = (search or "").strip()
        if search:
            like = f"%{search}%"
            query = query.filter(or_(Product.name.ilike(like), Product.barcode.ilike(like)))
        query = query.order_by(Product.name)

        status_filter = self.status_filter_combo.currentText() if hasattr(self, "status_filter_combo") else "الكل"

        if status_filter == "الكل":
            # ماكو فلتر حالة (يعتمد على حساب مخزون/صلاحية بالذاكرة) - نقدر
            # نطبّق الصفحة مباشرة على قاعدة البيانات (LIMIT/OFFSET)، فما نجيب
            # إلا الـ100 دواء المطلوبين فعليًا بهذي الصفحة، مو كل الأدوية.
            total_count = query.count()
            self._total_pages = max(1, -(-total_count // self.PAGE_SIZE))
            self._current_page = max(1, min(self._current_page, self._total_pages))
            products = (
                query.offset((self._current_page - 1) * self.PAGE_SIZE)
                .limit(self.PAGE_SIZE)
                .all()
            )
        else:
            # ⚠️ إصلاح أداء (راجع تقرير التدقيق، بند 14.5): كان هذا الفرع
            # يجيب *كل* الأدوية المطابقة لفلاتر الأرشفة/الاعتماد/البحث
            # ككائنات كاملة مع كل دفعاتها (query.all() بدون حد) فقط عشان
            # يحسب حالة كل دواء بايثونيًا - حتى لو الفلتر بالنهاية يرجّع
            # عدد قليل. الحين: (أ) نجيب فقط id + عتبة النقص لكل دواء مطابق
            # (بدون تحميل الدفعات)، (ب) نحسب مجموع المخزون/أقرب صلاحية لكل
            # دواء بتجميع SQL وحد (نفس فكرة إصلاح شارة التنبيهات
            # بـMainWindow._compute_alerts_count)، (ج) نطبّق نفس التصنيف
            # بايثونيًا على أرقام بسيطة بس، (د) نجيب الكائنات الكاملة (مع
            # batches) بس لمنتجات الصفحة الحالية اللي فعليًا رح تُعرض.
            id_query = self.session.query(Product.id, Product.min_stock_threshold)
            if not self.show_archived_check.isChecked():
                id_query = id_query.filter(Product.is_active == True)
            if self.approved_only_check.isChecked():
                id_query = id_query.filter(approved_clause(self.session))
            if search:
                like = f"%{search}%"
                id_query = id_query.filter(or_(Product.name.ilike(like), Product.barcode.ilike(like)))
            id_query = id_query.order_by(Product.name)
            candidates = id_query.all()
            candidate_ids = [pid for pid, _ in candidates]
            threshold_by_id = {pid: (thr or 10) for pid, thr in candidates}

            if status_filter == "راكد":
                stagnant_ids = self._stagnant_product_ids()
                matched_ids = [pid for pid in candidate_ids if pid in stagnant_ids]
            else:
                stats_by_id = self._batch_stats_for_ids(candidate_ids)
                today = date.today()
                alert_days = self._expiry_alert_days
                matched_ids = []
                for pid in candidate_ids:
                    stock, nearest = stats_by_id.get(pid, (0, None))
                    status_name = self._classify_status_name(stock, nearest, threshold_by_id[pid], today, alert_days)
                    if status_name == status_filter:
                        matched_ids.append(pid)

            total_count = len(matched_ids)
            self._total_pages = max(1, -(-total_count // self.PAGE_SIZE))
            self._current_page = max(1, min(self._current_page, self._total_pages))
            start = (self._current_page - 1) * self.PAGE_SIZE
            page_ids = matched_ids[start:start + self.PAGE_SIZE]

            if page_ids:
                # نحمّل الكائنات الكاملة (مع batches) بس لمنتجات هذي
                # الصفحة - IN() ما يضمن ترتيب النتائج، فنعيد ترتيبها
                # بايثونيًا حسب نفس تسلسل page_ids (مرتب بالاسم أصلًا من
                # id_query.order_by(Product.name) فوق).
                loaded = (
                    self.session.query(Product)
                    .options(joinedload(Product.batches))
                    .filter(Product.id.in_(page_ids))
                    .all()
                )
                by_id = {p.id: p for p in loaded}
                products = [by_id[pid] for pid in page_ids if pid in by_id]
            else:
                products = []

        self.page_info_label.setText(f"صفحة {self._current_page} من {self._total_pages} - إجمالي {total_count} دواء")
        self.prev_page_btn.setEnabled(self._current_page > 1)
        self.next_page_btn.setEnabled(self._current_page < self._total_pages)

        self.table.setRowCount(len(products))
        for row, p in enumerate(products):
            total_stock = sum(b.quantity_available for b in p.batches)
            nearest_expiry = min(
                (b.expiry_date for b in p.batches if b.expiry_date and b.quantity_available > 0),
                default=None,
            )
            status, status_color = self._status_for(p, total_stock, nearest_expiry)
            if not p.is_active:
                status = "مؤرشف"

            # سعر شراء الشريط والكمية بالباكيت محسوبة تلقائيًا من سعر شراء
            # الباكيت وعدد الأشرطة بداخله - نفس منطق الحساب المستخدم بنوافذ
            # إضافة/تعديل الدواء بالضبط، لعرضها هنا بدون تخزينها كأعمدة منفصلة.
            spc = p.strips_per_carton or 1
            strip_purchase_cost = (p.carton_purchase_price or 0) / spc if spc else 0
            qty_cartons = (total_stock / spc) if spc else total_stock

            values = [
                p.name, p.generic_name or "-", p.manufacturer or "-", p.category or "-", p.barcode or "-",
                f"{p.carton_purchase_price or 0:,.0f}",
                f"{p.carton_price or 0:,.0f}",
                f"{strip_purchase_cost:,.0f}",
                f"{p.sale_price:,.0f}",
                _fmt_qty(qty_cartons),
                str(total_stock),
                nearest_expiry.isoformat() if nearest_expiry else "-", status,
            ]
            # الأعمدة القابلة للتعديل مباشرة من الجدول: شراء/بيع الباكيت، بيع
            # الشريط، والكمية الفعلية بالشريط. سعر شراء الشريط والكمية بالباكيت
            # محسوبة تلقائيًا فقط (غير قابلة للتعديل مباشرة).
            editable_cols = {5, 6, 8, 10}
            for col, val in enumerate(values):
                item = QTableWidgetItem(val)
                if col in editable_cols and p.is_active:
                    item.setFlags(item.flags() | Qt.ItemIsEditable)
                else:
                    item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                if col == 12:
                    item.setForeground(QColor("#9CA3AF") if not p.is_active else status_color)
                if not p.is_active:
                    item.setBackground(QColor("#F1F5F9"))
                item.setData(Qt.UserRole, p.id)
                self.table.setItem(row, col, item)

            edit_btn = QPushButton("تعديل")
            edit_btn.setStyleSheet("background:#F1F5F9;border:1px solid #E5E7EB;border-radius:6px;padding:2px 6px;font-size:11px;")
            edit_btn.setAutoDefault(False)
            edit_btn.setDefault(False)
            edit_btn.clicked.connect(lambda _, pid=p.id: self.edit_product_by_id(pid))
            self.table.setCellWidget(row, 13, edit_btn)

            if p.is_active:
                action_btn = QPushButton("حذف")
                action_btn.setIcon(icon("trash", color="#DC2626", size=11))
                action_btn.setStyleSheet("background:#FEF2F2;border:1px solid #FECACA;color:#DC2626;border-radius:6px;padding:2px 6px;font-size:11px;")
                action_btn.clicked.connect(lambda _, pid=p.id: self.delete_product(pid))
            else:
                action_btn = QPushButton("إظهار")
                action_btn.setIcon(icon("undo", color="#15803D", size=11))
                action_btn.setStyleSheet("background:#DCFCE7;border:1px solid #DCFCE7;color:#15803D;border-radius:6px;padding:2px 6px;font-size:11px;")
                action_btn.clicked.connect(lambda _, pid=p.id: self.unarchive_product(pid))
            action_btn.setAutoDefault(False)
            action_btn.setDefault(False)
            self.table.setCellWidget(row, 14, action_btn)
        self.table.resizeRowsToContents()
        self._loading = False

    def _batch_stats_for_ids(self, product_ids):
        """مجموع الكمية المتاحة (كل الدفعات) + أقرب تاريخ صلاحية (لدفعات
        فيها كمية > 0 فقط) لكل منتج من القائمة - بتجميع SQL وحد (SUM/MIN
        مع GROUP BY)، بدل تحميل كل صفوف الدفعات ككائنات كاملة. نفس فكرة
        إصلاح شارة التنبيهات بـ MainWindow._compute_alerts_count بالضبط."""
        if not product_ids:
            return {}
        rows = (
            self.session.query(
                Batch.product_id,
                func.coalesce(func.sum(Batch.quantity_available), 0),
                func.min(
                    case(
                        (and_(Batch.expiry_date.isnot(None), Batch.quantity_available > 0), Batch.expiry_date),
                        else_=None,
                    )
                ),
            )
            .filter(Batch.product_id.in_(product_ids))
            .group_by(Batch.product_id)
            .all()
        )
        return {row[0]: (row[1], row[2]) for row in rows}

    @staticmethod
    def _classify_status_name(stock, nearest_expiry, threshold, today, alert_days):
        """نفس منطق التصنيف بـ_status_for بالضبط، بس على أرقام بسيطة
        (مخزون/أقرب صلاحية/عتبة) بدون الحاجة لكائن Product كامل محمّل -
        يسمح بتطبيق نفس التصنيف على نتائج استعلام تجميع SQL مباشرة."""
        if stock <= 0:
            return "نقص"
        if nearest_expiry and nearest_expiry <= today:
            return "منتهي"
        if nearest_expiry and nearest_expiry <= today + timedelta(days=alert_days):
            return "قريب الانتهاء"
        if stock <= threshold:
            return "منخفض"
        return "OK"

    _STATUS_COLORS = {
        "نقص": "#DC2626",
        "منتهي": "#DC2626",
        "قريب الانتهاء": "#F59E0B",
        "منخفض": "#F59E0B",
        "OK": "#16A34A",
    }

    def _status_for(self, product, stock, nearest_expiry):
        from PySide6.QtGui import QColor
        today = date.today()
        threshold = product.min_stock_threshold or 10
        alert_days = getattr(self, "_expiry_alert_days", None) or get_expiry_alert_days(self.session)
        name = self._classify_status_name(stock, nearest_expiry, threshold, today, alert_days)
        return name, QColor(self._STATUS_COLORS[name])

    def _on_item_changed(self, item):
        if self._loading:
            return
        product_id = item.data(Qt.UserRole)
        product = self.session.query(Product).get(product_id)
        if not product:
            return
        col = item.column()
        try:
            if col == 5:
                product.carton_purchase_price = float(item.text())
            elif col == 6:
                product.carton_price = float(item.text())
            elif col == 8:
                product.sale_price = float(item.text())
            elif col == 10 and product.batches:  # الكمية على آخر دفعة (تبسيط)
                product.batches[-1].quantity_available = int(item.text())
            self.session.commit()
            # نحدّث الجدول عشان الأعمدة المحسوبة تلقائيًا (سعر شراء الشريط
            # والكمية بالباكيت) تنعكس فورًا لو غيّرت سعر/كمية الباكيت.
            if col in (5, 10):
                self.refresh(self.barcode_scan_input.text())
        except ValueError:
            QMessageBox.warning(self, "خطأ", "القيمة المدخلة غير صحيحة.")
            self.refresh()

    def add_product(self):
        dialog = AddProductDialog(self)
        self._save_new_product(dialog)

    def _save_new_product(self, dialog):
        dialog.raise_()
        dialog.activateWindow()
        if dialog.exec() == QDialog.Accepted:
            data = dialog.get_data()
            if not data["name"]:
                QMessageBox.warning(self, "تنبيه", "لازم تدخل اسم الدواء.")
                return
            product = Product(
                name=data["name"], generic_name=data["generic_name"],
                manufacturer=data["manufacturer"], category=data["category"],
                barcode=data["barcode"], base_unit="شريط", is_custom=True,
                sale_price=data["sale_price"], wholesale_price=data["wholesale_price"],
                carton_price=data["carton_price"], min_stock_threshold=data["min_threshold"],
                strips_per_carton=data["strips_per_carton"], carton_purchase_price=data["carton_cost"],
            )
            self.session.add(product)
            self.session.flush()
            for unit_name, factor in [("شريط", 1), ("علبة", data["strips_per_carton"])]:
                self.session.add(ProductUnit(product_id=product.id, unit_name=unit_name, conversion_factor=factor))
            if data["qty"] > 0:
                self.session.add(Batch(
                    product_id=product.id, batch_number="MANUAL-0001",
                    expiry_date=data["expiry"], purchase_price=data["cost"],
                    quantity_received=data["qty"], quantity_available=data["qty"],
                    carton_purchase_price=data["carton_cost"], carton_strips_per_carton=data["strips_per_carton"],
                ))
            self.session.commit()
            self.refresh()

    def delete_product(self, product_id):
        product = self.session.query(Product).get(product_id)
        if not product:
            return
        was_sold = (
            self.session.query(InvoiceItem)
            .join(Batch, InvoiceItem.batch_id == Batch.id)
            .filter(Batch.product_id == product.id)
            .first()
            is not None
        )
        if was_sold:
            reply = QMessageBox.question(
                self, "أرشفة الدواء",
                f"الدواء \"{product.name}\" بيع منه سابقًا بفواتير، فما نقدر نحذفه نهائيًا حتى ما تنكسر فواتيرك وتقاريرك القديمة.\n\n"
                "بس نقدر نؤرشفه (يختفي من المخزون والبيع، وتقدر ترجعه أي وقت من زر \"إظهار المؤرشفة\").\n\nتريد تأرشفه هسه؟",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No
            )
            if reply != QMessageBox.Yes:
                return
            product.is_active = False
            self.session.commit()
            QMessageBox.information(self, "تم", f"تم أرشفة \"{product.name}\" وإخفاؤه من المخزون والبيع.")
        else:
            reply = QMessageBox.question(
                self, "حذف نهائي",
                f"متأكد تريد تحذف \"{product.name}\" نهائيًا؟ هذا الدواء ما بيع منه أي مرة، فالحذف كامل ولا يمكن التراجع عنه.",
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No
            )
            if reply != QMessageBox.Yes:
                return
            self.session.delete(product)
            self.session.commit()
            QMessageBox.information(self, "تم", "تم حذف الدواء نهائيًا.")
        self.refresh(self.barcode_scan_input.text())

    def unarchive_product(self, product_id):
        product = self.session.query(Product).get(product_id)
        if not product:
            return
        product.is_active = True
        self.session.commit()
        QMessageBox.information(self, "تم", f"تم إرجاع \"{product.name}\" للمخزون والبيع.")
        self.refresh(self.barcode_scan_input.text())

    def edit_product_by_id(self, product_id):
        product = self.session.query(Product).get(product_id)
        if product:
            self.edit_product(product)

    def edit_product(self, product):
        dialog = EditProductDialog(product, self)
        dialog.raise_()
        dialog.activateWindow()
        if dialog.exec() == QDialog.Accepted:
            data = dialog.get_data()
            if not data["name"]:
                QMessageBox.warning(self, "تنبيه", "لازم تدخل اسم الدواء.")
                return
            product.name = data["name"]
            product.generic_name = data["generic_name"]
            product.manufacturer = data["manufacturer"]
            product.category = data["category"]
            product.barcode = data["barcode"]
            product.sale_price = data["sale_price"]
            product.carton_price = data["carton_price"]
            product.carton_purchase_price = data["carton_purchase_price"]
            if data.get("update_existing_batches"):
                # نستخدم عدد الأشرطة الجديد (data["strips_per_carton"]) اللي
                # المستخدم عدله بنفس الشاشة - مو product.strips_per_carton
                # القديم لسا (يتحدث بالسطر تحت)، وإلا يطلع سعر الشريط المحسوب
                # غلط لو المستخدم غيّر عدد الأشرطة وأشر "حدّث كل الدفعات" بنفس الوقت.
                spc = data["strips_per_carton"] or 1
                new_strip_cost = (data["carton_purchase_price"] / spc) if spc else 0
                for b in product.batches:
                    if b.quantity_available > 0 and round(b.purchase_price or 0, 2) != round(new_strip_cost, 2):
                        old_price = b.purchase_price or 0
                        b.purchase_price = new_strip_cost
                        b.carton_purchase_price = data["carton_purchase_price"]
                        b.carton_strips_per_carton = spc
                        self.session.add(StockMovement(
                            batch_id=b.id, movement_type="تصحيح سعر", quantity=0,
                            note=f"تصحيح سعر شراء الشريط من {old_price:,.0f} إلى {new_strip_cost:,.0f} (تحديث كل الدفعات)",
                        ))
            product.min_stock_threshold = data["min_threshold"]
            if product.strips_per_carton != data["strips_per_carton"]:
                product.strips_per_carton = data["strips_per_carton"]
                carton_unit = next((u for u in product.units if u.unit_name != "شريط"), None)
                if carton_unit:
                    carton_unit.conversion_factor = data["strips_per_carton"]
            for batch_id, new_expiry, new_qty, batch_spc in data.get("batch_updates", []):
                batch = self.session.query(Batch).get(batch_id)
                if batch:
                    batch.expiry_date = new_expiry
                    if new_qty != batch.quantity_available and new_qty >= 0:
                        diff = new_qty - batch.quantity_available
                        self.session.add(StockMovement(
                            batch_id=batch.id, movement_type="تصحيح يدوي", quantity=abs(diff),
                            note=f"تصحيح كمية من {batch.quantity_available} إلى {new_qty} من شاشة تعديل الدواء",
                        ))
                        batch.quantity_available = new_qty
                    # نثبّت عدد الأشرطة الخاص بهذي الدفعة لو ماكان مسجّل من
                    # قبل (دفعة قديمة) عشان تصير التحويلات الجاية لهذي الدفعة صحيحة.
                    if not batch.carton_strips_per_carton:
                        batch.carton_strips_per_carton = batch_spc
            if data["add_qty"] > 0:
                # نسجّل فاتورة شراء فعلية بالمورد المحدد (لو أكو) - هذا اللي
                # يخلي الدفعة تظهر بكشف حساب المورد بالموردين، مو دفعة
                # "يتيمة" بدون مصدر معروف زي ما كان يصير قبل هذا التعديل.
                order = PurchaseOrder(
                    supplier_id=data["add_supplier_id"], status="confirmed", source="شاشة المخزون",
                    total_amount=data["add_carton_cost"] * (data["add_qty"] / data["add_spc"] if data["add_spc"] else 0),
                    receipt_number=data["add_receipt_number"], order_date=data["add_receipt_date"],
                )
                self.session.add(order)
                self.session.flush()
                self.session.add(PurchaseOrderItem(
                    purchase_order_id=order.id, product_id=product.id,
                    quantity=data["add_qty"] / data["add_spc"] if data["add_spc"] else 0,
                    unit_cost=data["add_carton_cost"],
                ))
                self.session.add(Batch(
                    product_id=product.id, batch_number=f"PO-{order.id}-{product.id}",
                    expiry_date=data["add_expiry"], purchase_price=data["add_cost"],
                    quantity_received=data["add_qty"], quantity_available=data["add_qty"],
                    carton_purchase_price=data["add_carton_cost"], carton_strips_per_carton=data["add_spc"],
                    supplier_id=data["add_supplier_id"], receipt_number=data["add_receipt_number"],
                ))
            if data["waste_qty"] > 0:
                remaining = data["waste_qty"]
                batches = sorted(
                    [b for b in product.batches if b.quantity_available > 0],
                    key=lambda b: (b.expiry_date or date.max)
                )
                for batch in batches:
                    if remaining <= 0:
                        break
                    take = min(batch.quantity_available, remaining)
                    batch.quantity_available -= take
                    remaining -= take
                    self.session.add(StockMovement(
                        batch_id=batch.id, movement_type="تالف", quantity=take,
                        note=data["waste_reason"] or "هدر غير محدد السبب",
                    ))
            self.session.commit()
            self.refresh()

    def export_excel(self):
        try:
            from openpyxl import Workbook
        except ImportError:
            QMessageBox.warning(self, "مكتبة ناقصة", "لازم تنصب openpyxl: pip install openpyxl")
            return
        path, _ = QFileDialog.getSaveFileName(self, "تصدير المخزون", "المخزون.xlsx", "Excel Files (*.xlsx)")
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.append(["الاسم التجاري", "الاسم العلمي", "الشركة", "التصنيف", "الباركود", "المخزون",
                   "سعر الشراء الباكيت", "سعر الشراء (شريط)", "سعر البيع باكيت", "سعر البيع (شريط)", "التعبئة"])
        for p in self.session.query(Product).filter(Product.is_active == True).all():
            stock = sum(b.quantity_available for b in p.batches)
            spc = p.strips_per_carton or 1
            strip_purchase_price = (p.carton_purchase_price or 0) / spc if spc else 0
            ws.append([p.name, p.generic_name, p.manufacturer, p.category, p.barcode,
                       stock, p.carton_purchase_price, strip_purchase_price, p.carton_price, p.sale_price, p.strips_per_carton])
        wb.save(path)
        QMessageBox.information(self, "تم", "تم تصدير المخزون بنجاح.")

    def import_excel(self):
        """
        يستورد أدوية من ملف إكسل بنفس ترتيب الأعمدة تبع التصدير:
        الاسم التجاري، الاسم العلمي، الشركة، التصنيف، الباركود، المخزون،
        سعر الشراء الباكيت، سعر الشراء (شريط)، سعر البيع باكيت، سعر البيع (شريط)، التعبئة

        ملاحظة: "سعر الشراء (شريط)" عمود معلوماتي بس بالتصدير (محسوب تلقائيًا
        = سعر شراء الباكيت ÷ التعبئة) - عند الاستيراد نتجاهل قيمته حتى لو
        كانت موجودة بالملف ونعيد حسابه من سعر الباكيت والتعبئة دائمًا.

        ملاحظة أداء/موثوقية مهمة (مع ملفات كبيرة - آلاف الأدوية): الاستيراد
        الفعلي يشتغل بخيط منفصل (_ImportExcelWorker) عن خيط الواجهة - قبل
        هذا الإصلاح كانت الواجهة تتجمّد تمامًا (ويندوز يعرضها "لا يستجيب")
        طول مدة الاستيراد، رغم إنها كانت شغّالة فعليًا بالخلفية. الحين
        الواجهة تضل تستجيب، وفيه شريط تقدم حقيقي (صف X من Y) + زر إلغاء.
        كل صف يُعالج بـ try/except + SAVEPOINT مستقل (صف فيه خلية غير نظيفة
        يُتخطى بدل ما يوقف كل شي)، مع حفظ دوري كل 500 صف حتى ما يضيع التقدم
        لو صار أي عطل لاحق.
        """
        path, _ = QFileDialog.getOpenFileName(self, "استيراد أدوية", "", "Excel Files (*.xlsx)")
        if not path:
            return

        self._import_progress = QProgressDialog("جاري تجهيز الملف...", "إلغاء", 0, 0, self)
        self._import_progress.setWindowTitle("استيراد أدوية")
        self._import_progress.setWindowModality(Qt.WindowModal)
        self._import_progress.setMinimumDuration(0)
        self._import_progress.setAutoClose(False)
        self._import_progress.setAutoReset(False)

        self._import_worker = _ImportExcelWorker(path)
        self._import_worker.progress.connect(self._on_import_progress)
        self._import_worker.finished_ok.connect(self._on_import_finished)
        self._import_worker.finished_error.connect(self._on_import_error)
        self._import_progress.canceled.connect(self._import_worker.cancel)
        self._import_worker.start()
        self._import_progress.show()

    def _on_import_progress(self, current, total):
        if total:
            self._import_progress.setRange(0, total)
            self._import_progress.setValue(current)
            self._import_progress.setLabelText(f"جاري الاستيراد... {current} من {total}")
        else:
            self._import_progress.setLabelText(f"جاري الاستيراد... صف {current}")

    def _on_import_error(self, msg):
        self._import_progress.close()
        QMessageBox.warning(self, "خطأ", msg)

    def _on_import_finished(self, result):
        self._import_progress.close()
        self.refresh()
        count = result["count"]
        inserted = result["inserted"]
        updated = result["updated"]
        total_rows_seen = result["total_rows_seen"]
        errors = result["errors"]
        commit_failed = result["commit_failed"]
        cancelled = result["cancelled"]

        msg = (
            f"عدد صفوف الملف اللي فيها بيانات: {total_rows_seen}\n"
            f"تمت معالجتها بنجاح: {count} ({inserted} دواء جديد، {updated} تحديث لدواء موجود)"
        )
        if cancelled:
            msg = "تم إلغاء الاستيراد.\n\n" + msg + "\n(التقدم لين لحظة الإلغاء محفوظ فعليًا)"
        if commit_failed:
            msg += "\n\n⚠️ توقف الاستيراد بسبب فشل بالحفظ (التفاصيل بالأسفل) - راجع المشكلة قبل ما تعيد المحاولة."
        if errors:
            preview = "\n".join(errors[:15])
            more = f"\n... و{len(errors) - 15} صف إضافي" if len(errors) > 15 else ""
            msg += f"\n\nتعذّر استيراد {len(errors)} صف:\n{preview}{more}"
        QMessageBox.information(self, "تم", msg)
