"""تعبئة قاعدة البيانات بقائمة الأدوية الأولية.

هذا الملف يشتغل مرة وحدة بس (أول تشغيل على قاعدة بيانات فاضية تمامًا -
راجع الشرط `if session.query(Product).count() > 0` تحت). المصدر: ملف
"starter_catalog.xlsx" بنفس مجلد هذا الملف (app/db/) - كتالوج حقيقي
مرفوع من صاحب البرنامج (11,082 صنف)، محل القائمة الثابتة القديمة الـ85
صنف.

مهم - هذا الملف قابل للتحديث بنفسك وقت ما تريد:
افتح app/db/starter_catalog.xlsx بإكسل، عدّل/ضيف/احذف صفوف براحتك، واحفظه
بنفس الاسم والمكان ونفس ترتيب الأعمدة (الاسم التجاري، الاسم العلمي،
الشركة، التصنيف، الباركود، المخزون، سعر شراء الباكيت، سعر شراء الشريط،
سعر بيع الباكيت، سعر بيع الشريط، التعبئة). أي تنصيب جديد أو قاعدة بيانات
فاضية بالمستقبل تاخذ القائمة المحدّثة تلقائيًا من نفس الملف.

ملاحظة: النسخة القديمة من هذا الملف كانت تزرع كمان أمثلة تجريبية لإرشادات
جرعات وتداخلات دوائية وبدائل، مربوطة بأسماء أدوية القائمة الـ85 القديمة
بالضبط. بما إن القائمة الجديدة كتالوج حقيقي مختلف الأسماء بالكامل، انشالت
هذي الأمثلة (كانت أصلاً معلّمة "تجريبي - راجعها قبل الاعتماد" ومو بيانات
حقيقية) - أضف إرشادات الجرعات والتداخلات الدوائية الحقيقية بنفسك من شاشاتها
المخصصة بالبرنامج وقت ما تحتاج.
"""
import os
import re

from .database import get_session, init_db
from .models import Product, ProductUnit, Batch

CATALOG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "starter_catalog.xlsx")

BASE_UNIT_NAME = "شريط"
CARTON_UNIT_NAME = "علبة"

# ترتيب أعمدة ملف الكتالوج (starter_catalog.xlsx) - بالضبط بهذا الترتيب.
COLUMNS = [
    "name", "generic", "manufacturer", "category", "barcode",
    "stock", "carton_cost", "strip_cost", "carton_price", "sale_price", "packaging",
]


def _text(raw):
    if raw is None:
        return ""
    return str(raw).strip()


def _num(raw):
    """يحوّل أي قيمة خام (رقم، نص، فاضي) لعدد عشري - 0 لو ما قدر يفهمها."""
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    try:
        return float(str(raw).strip())
    except ValueError:
        return 0.0


def _parse_packaging(raw):
    """عمود "التعبئة": مثلًا "3 شريط/علبة" -> 3. فاضي/غير مفهوم -> 1
    (يعني نعتبره منتج "علبة وحدة بس" بدون أشرطة، بنفس منطق نافذة إضافة
    دواء بسطح المكتب والموبايل)."""
    match = re.match(r"\s*(\d+)", _text(raw))
    if match:
        value = int(match.group(1))
        if value > 0:
            return value
    return 1


def _parse_stock(raw, strips_per_carton):
    """عمود "المخزون": رقم صافي (بدون وحدة مكتوبة) يُعتبر أصلاً بوحدة
    الشريط مباشرة (الصيغة الغالبة بالملف). القيم النصية المكتوب معها
    "علبة" (مثل "5 علبة") تتحول لشريط بضربها بعدد الأشرطة بالعلبة."""
    if raw is None:
        return 0
    if isinstance(raw, (int, float)):
        return max(int(raw), 0)
    text = _text(raw)
    match = re.match(r"\s*(\d+)", text)
    if not match:
        return 0
    value = int(match.group(1))
    if "علبة" in text:
        value = value * strips_per_carton
    return max(value, 0)


def _iter_catalog_rows():
    from openpyxl import load_workbook
    wb = load_workbook(CATALOG_FILE, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not _text(row[0]):
            continue
        values = list(row) + [None] * (len(COLUMNS) - len(row))
        yield dict(zip(COLUMNS, values[: len(COLUMNS)]))


def _import_catalog_rows(session, skip_existing_names=False):
    """يقرأ starter_catalog.xlsx ويضيف كل صف كـ Product جديد (+ وحداته +
    دفعته لو عنده مخزون حقيقي). مستخدمة من seed() (أول تشغيل بس) ومن
    dev_tools/import_catalog_now.py (استيراد فوري لقاعدة بيانات موجودة
    مسبقًا). يرجّع عدد الأصناف اللي انضافت فعلاً."""
    existing_names = set()
    if skip_existing_names:
        existing_names = {n for (n,) in session.query(Product.name).all()}

    added = 0
    for row in _iter_catalog_rows():
        name = _text(row["name"])
        if not name or name in existing_names:
            continue

        strips_per_carton = _parse_packaging(row["packaging"])
        carton_purchase_price = _num(row["carton_cost"])
        carton_price = _num(row["carton_price"])
        sale_price = _num(row["sale_price"])
        strip_cost = _num(row["strip_cost"]) or (
            (carton_purchase_price / strips_per_carton) if strips_per_carton else 0
        )
        # نفس منطق نافذة "إضافة دواء جديد": لو المنتج "علبة وحدة بس"
        # (شريط واحد بالعلبة)، نوحّد سعر الشريط مع سعر العلبة - أي وحد
        # فيهم موجود بالملف يعبّي الثاني.
        if strips_per_carton == 1:
            sale_price = sale_price or carton_price
            carton_price = carton_price or sale_price
        qty = _parse_stock(row["stock"], strips_per_carton)

        product = Product(
            name=name, generic_name=_text(row["generic"]), manufacturer=_text(row["manufacturer"]),
            category=_text(row["category"]), base_unit=BASE_UNIT_NAME, is_custom=False,
            barcode=_text(row["barcode"]) or None,
            sale_price=sale_price, wholesale_price=0, carton_price=carton_price,
            strips_per_carton=strips_per_carton, carton_purchase_price=carton_purchase_price,
        )
        session.add(product)
        session.flush()  # يعطينا product.id قبل الـ commit
        added += 1
        if skip_existing_names:
            existing_names.add(name)

        for unit_name, factor in [(BASE_UNIT_NAME, 1), (CARTON_UNIT_NAME, strips_per_carton)]:
            session.add(ProductUnit(product_id=product.id, unit_name=unit_name, conversion_factor=factor))

        # دفعة فعلية بس لو الكتالوج فيه مخزون حقيقي أكبر من صفر لهذا
        # الصنف - ما نخترع مخزون وهمي لصنف مكتوب مخزونه 0 بالملف.
        if qty > 0:
            session.add(Batch(
                product_id=product.id, batch_number=f"CATALOG-{product.id}",
                expiry_date=None,
                purchase_price=strip_cost, quantity_received=qty, quantity_available=qty,
            ))

        if added % 500 == 0:
            session.flush()

    session.commit()
    return added


def seed():
    init_db()
    session = get_session()

    if session.query(Product).count() > 0:
        print("قاعدة البيانات معبّاة مسبقًا - تخطي التعبئة.")
        return  # ملاحظة: ما نسكر الجلسة المشتركة (scoped_session) هنا أبدًا

    if not os.path.exists(CATALOG_FILE):
        print(f"⚠️ ملف الكتالوج {CATALOG_FILE} غير موجود - تخطي التعبئة الأولية.")
        return

    added = _import_catalog_rows(session, skip_existing_names=False)
    print(f"تمت تعبئة {added} دواء من كتالوج {os.path.basename(CATALOG_FILE)} بنجاح.")


if __name__ == "__main__":
    seed()
